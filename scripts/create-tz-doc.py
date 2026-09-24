#!/usr/bin/env python3
"""Generate contractor TZ as a Word document."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()

# Styles
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(26, 26, 46)
    return h

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
    # Rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    return table

# ── Title ──
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('ТЕХНИЧЕСКОЕ ЗАДАНИЕ')
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(26, 26, 46)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Печать и доставка картин для проекта Шежіре (skezire.kz)')
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph()

# ── 1. Description ──
add_heading('1. Описание проекта', level=2)
doc.add_paragraph(
    'Сайт skezire.kz — казахское генеалогическое дерево с AI-функциями. '
    'Пользователи генерируют AI-фото (портрет 100 лет назад, экшн-фигурка, стиль Гибли и др.) '
    'и могут заказать печать результата на холсте/постере с доставкой по Казахстану.'
)

# ── 2. Catalog ──
add_heading('2. Каталог продукции', level=2)
add_table(
    ['Формат', 'Размер', 'Цена для клиента (KZT)'],
    [
        ['Постер', 'A3', '3 990'],
        ['Холст', '30x40 см', '7 990'],
        ['Холст', '50x70 см', '12 990'],
        ['Холст в рамке', '30x40 см', '11 990'],
        ['Холст в рамке', '50x70 см', '17 990'],
    ]
)
doc.add_paragraph()
doc.add_paragraph(
    'Доставка: бесплатная для клиента (включена в цену). '
    'Расходы на доставку — на стороне подрядчика или оговариваются отдельно.'
)

# ── 3. How orders arrive ──
add_heading('3. Как подрядчик получает заказы', level=2)
doc.add_paragraph(
    'Заказы приходят через Telegram-бот автоматически. '
    'Подрядчику предоставляется доступ в админ-чат бота.'
)

add_heading('Что приходит в уведомлении о заказе:', level=3)
items = [
    'Номер заказа (например, #1001)',
    'Товар: тип + размер (напр. "Холст 50x70 см")',
    'Сумма заказа',
    'Статус оплаты (заказы приходят только после оплаты)',
    'Данные получателя: ФИО, телефон, город, адрес',
    'Ссылка на изображение для печати (оригинал + резервная копия)',
    'Тип AI-генерации (для понимания формата)',
]
for item in items:
    doc.add_paragraph(item, style='List Bullet')

add_heading('Команды бота:', level=3)
add_table(
    ['Команда', 'Что делает'],
    [
        ['/orders', 'Список активных заказов (оплаченные, в работе, отправленные)'],
        ['/order 1001', 'Полная информация по заказу #1001'],
        ['/track {ID} {ТРЕК}', 'Добавить трек-номер и отметить отправку'],
        ['/stats', 'Статистика: кол-во заказов, выручка'],
    ]
)

doc.add_paragraph()
add_heading('Кнопки управления (в каждом заказе):', level=3)
buttons = [
    '"В работу" — подрядчик нажимает, когда начинает печать',
    '"Отправлено" — бот просит ввести трек-номер',
    '"Доставлено" — после подтверждения получения клиентом',
]
for b in buttons:
    doc.add_paragraph(b, style='List Bullet')

# ── 4. Process ──
add_heading('4. Процесс выполнения заказа', level=2)
p = doc.add_paragraph()
run = p.add_run('Оплачен  →  В работу  →  Отправлен  →  Доставлен')
run.bold = True
run.font.size = Pt(12)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
add_table(
    ['Этап', 'Действие подрядчика', 'Срок'],
    [
        ['1. Получение заказа', 'Проверить уведомление в Telegram, скачать изображение', '—'],
        ['2. Связь с клиентом', 'Позвонить/написать на указанный телефон, подтвердить адрес', 'В течение 24 ч'],
        ['3. Начало печати', 'Нажать "В работу" в боте', '—'],
        ['4. Печать', 'Распечатать изображение в заказанном формате', '1-3 рабочих дня'],
        ['5. Упаковка', 'Надёжная упаковка (тубус для постеров, коробка для холстов)', '—'],
        ['6. Отправка', 'Отправить через курьерскую службу, ввести трек-номер через /track', '—'],
        ['7. Контроль доставки', 'Отслеживать получение, при подтверждении нажать "Доставлено"', '—'],
    ]
)

# ── 5. Print requirements ──
add_heading('5. Требования к печати', level=2)
reqs = [
    'Изображения: AI-генерации разрешением ~1024x1024 до ~1536x1024 px. Формат: JPG/PNG.',
    'Качество: При необходимости — апскейл до нужного разрешения перед печатью.',
    'Постер A3: Матовая/глянцевая бумага, плотность от 200 г/м\u00b2.',
    'Холст: Натяжка на подрамник, галерейная намотка (края загнуты).',
    'Холст в рамке: Подрамник + рамка (цвет: чёрный или натуральное дерево — уточнить с клиентом).',
]
for r in reqs:
    doc.add_paragraph(r, style='List Bullet')

# ── 6. Packaging & delivery ──
add_heading('6. Требования к упаковке и доставке', level=2)
pkg = [
    'Постеры — в жёстком тубусе.',
    'Холсты — в картонной коробке с уголками.',
    'Доставка по всему Казахстану.',
    'Трек-номер обязателен для каждой отправки.',
    'Срок доставки клиенту: до 7 рабочих дней (печать + доставка).',
]
for p in pkg:
    doc.add_paragraph(p, style='List Bullet')

# ── 7. Client communication ──
add_heading('7. Связь с клиентом', level=2)
doc.add_paragraph('Подрядчик обязан связаться с клиентом после получения заказа для:')
comms = [
    'Подтверждения адреса доставки',
    'Уточнения деталей (цвет рамки, если применимо)',
    'Уведомления об отправке и трек-номере',
    'Решения проблем (брак, повреждение при доставке)',
]
for c in comms:
    doc.add_paragraph(c, style='List Bullet')
doc.add_paragraph('Контакт клиента (телефон) указан в каждом заказе.')

# ── 8. Financial ──
add_heading('8. Финансовые условия', level=2)

add_heading('Вариант А — Фиксированная цена за единицу:', level=3)
add_table(
    ['Формат', 'Цена подрядчика (KZT)', 'Маржа сайта'],
    [
        ['Постер A3', '___________', '___________'],
        ['Холст 30x40', '___________', '___________'],
        ['Холст 50x70', '___________', '___________'],
        ['Холст в рамке 30x40', '___________', '___________'],
        ['Холст в рамке 50x70', '___________', '___________'],
    ]
)
doc.add_paragraph()
p = doc.add_paragraph()
run = p.add_run('* Включая печать + упаковку + доставку.')
run.italic = True
run.font.size = Pt(9)

add_heading('Вариант Б — Процент от заказа:', level=3)
doc.add_paragraph(
    'Подрядчик получает X% от суммы заказа за печать + упаковку + доставку.'
)

doc.add_paragraph()
doc.add_paragraph(
    'Оплата подрядчику: по факту выполнения (после статуса "Отправлен"), '
    'раз в неделю / раз в месяц — по договорённости.'
)

# ── 9. SLA ──
add_heading('9. SLA (уровень сервиса)', level=2)
add_table(
    ['Метрика', 'Требование'],
    [
        ['Время реакции на заказ', 'До 24 часов'],
        ['Время печати', 'До 3 рабочих дней'],
        ['Общий срок (заказ -> отправка)', 'До 4 рабочих дней'],
        ['Процент брака', 'Не более 3%'],
        ['Возврат при браке', 'Перепечатка за счёт подрядчика'],
    ]
)

# ── 10. What site provides ──
add_heading('10. Что предоставляет сайт', level=2)
site_provides = [
    'Автоматические уведомления о новых заказах через Telegram-бот',
    'Изображения для печати в максимальном качестве',
    'Данные клиента (ФИО, телефон, адрес)',
    'Систему отслеживания статусов',
    'Оплату от клиента (RoboCassa) — подрядчик не принимает деньги от клиентов',
]
for s in site_provides:
    doc.add_paragraph(s, style='List Bullet')

# ── 11. What contractor provides ──
add_heading('11. Что предоставляет подрядчик', level=2)
contractor = [
    'Оборудование и материалы для печати',
    'Упаковочные материалы',
    'Организацию доставки по Казахстану',
    'Связь с клиентом',
    'Гарантию качества печати',
]
for c in contractor:
    doc.add_paragraph(c, style='List Bullet')

# ── Footer ──
doc.add_paragraph()
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Просьба заполнить цены в разделе 8 и прислать КП.')
run.italic = True
run.font.color.rgb = RGBColor(100, 100, 100)

# Save
output_path = os.path.expanduser('~/Projects/skezire/TZ_Podryadchik_Skezire.docx')
doc.save(output_path)
print(f'Saved: {output_path}')
